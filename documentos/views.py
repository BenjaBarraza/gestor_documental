from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from documentos.models import Documento
from .forms import DocumentoForm, EnlacePublicoForm
from django.db.models import Q
from django.http import JsonResponse
from django.contrib import messages
import mimetypes
import docx
import openpyxl
from django.utils.timezone import now
from django.views.decorators.http import require_POST
from django.core.mail import EmailMultiAlternatives
from django.template.loader import render_to_string 
import uuid




def home(request):
    return render(request, 'documentos/home.html')

@login_required
def subir_documento(request):
    if request.method == 'POST':
        form = DocumentoForm(request.POST, request.FILES)
        if form.is_valid():
            documento = form.save(commit=False)
            documento.usuario = request.user
            documento.save()

            if (form.cleaned_data['archivo'].name.lower().endswith('.jpg') and
                request.FILES['archivo'].name.lower().endswith('.heic')):
                messages.info(request, "El archivo .HEIC fue convertido a .JPG automáticamente.")

            return redirect('documentos:lista')
    else:
        form = DocumentoForm()
    return render(request, 'documentos/subir.html', {'form': form})

@login_required
def listar_documentos(request):
    documentos = Documento.objects.all() if request.user.is_superuser else Documento.objects.filter(usuario=request.user)
    for doc in documentos:
        doc.etiquetas_lista = [e.strip() for e in doc.etiquetas.split(',')] if doc.etiquetas else []
    return render(request, 'documentos/lista.html', {'documentos': documentos})

@login_required
def buscar_documentos(request):
    query = request.GET.get("q", "").strip()
    documentos = Documento.objects.none()
    if query:
        documentos = Documento.objects.filter(Q(nombre__icontains=query) | Q(etiquetas__icontains=query) | Q(archivo__iendswith=query))
        if not request.user.is_superuser:
            documentos = documentos.filter(usuario=request.user)
    return render(request, 'documentos/resultados_busqueda.html', {'documentos': documentos, 'query': query})

@login_required
def previsualizar_documento(request, documento_id):
    documento = get_object_or_404(Documento, id=documento_id)
    if not request.user.is_superuser and documento.usuario != request.user:
        return render(request, 'documentos/403.html', status=403)

    mimetype, _ = mimetypes.guess_type(documento.archivo.url)
    es_pdf = mimetype == 'application/pdf'
    es_imagen = mimetype and mimetype.startswith('image/')
    es_video = mimetype and mimetype.startswith('video/')
    es_docx = mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    es_xlsx = mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    archivo_url_absoluto = request.build_absolute_uri(documento.archivo.url)

    contenido_docx = [p.text for p in docx.Document(documento.archivo.path).paragraphs if p.text.strip()] if es_docx else None
    contenido_xlsx = list(openpyxl.load_workbook(documento.archivo.path).active.iter_rows(values_only=True)) if es_xlsx else None

    return render(request, 'documentos/previsualizar.html', {
        'documento': documento,
        'es_pdf': es_pdf,
        'es_imagen': es_imagen,
        'es_video': es_video,
        'es_docx': es_docx,
        'es_xlsx': es_xlsx,
        'contenido_docx': contenido_docx,
        'contenido_xlsx': contenido_xlsx,
        'archivo_url_absoluto': archivo_url_absoluto,
    })

@require_POST
@login_required
def eliminar_documento(request, documento_id):
    documento = get_object_or_404(Documento, id=documento_id)
    if not request.user.is_superuser and documento.usuario != request.user:
        return JsonResponse({'status': 'error', 'message': 'Acceso denegado'}, status=403)
    try:
        documento.archivo.delete()
        documento.delete()
        return JsonResponse({'status': 'success'})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)
    

@login_required
def generar_enlace(request, doc_id):
    documento = get_object_or_404(Documento, id=doc_id)

    # Verificación de permisos
    if not request.user.is_superuser and documento.usuario != request.user:
        if request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'status': 'error', 'message': 'Acceso denegado'}, status=403)
        else:
            return render(request, 'documentos/prohibido.html', status=403)

    if request.method == 'POST':
        form = EnlacePublicoForm(request.POST)
        if form.is_valid():
            # Guardar expiración y enlace único
            documento.fecha_expiracion = form.cleaned_data['fecha_expiracion']
            documento.enlace_publico = str(uuid.uuid4())
            documento.save()

            # Construir enlace absoluto
            enlace_url = request.build_absolute_uri(f"/publico/{documento.enlace_publico}")

            # Guardar en sesión para mostrar tras redirect
            request.session['enlace_generado'] = enlace_url

            # ✉️ Enviar correo al usuario
            try:
                html_content = render_to_string("emails/enlace_generado.html", {
                    'usuario': request.user,
                    'documento': documento,
                    'enlace_url': enlace_url,
                })

                email = EmailMultiAlternatives(
                    subject='🔗 Enlace público generado',
                    body='Tu cliente de correo no admite HTML.',
                    from_email=None,  # usa DEFAULT_FROM_EMAIL
                    to=[request.user.email],
                )
                email.attach_alternative(html_content, "text/html")
                email.send()
                print(f"✅ Correo enviado a: {request.user.email}")
            except Exception as e:
                print("❌ ERROR al enviar correo:", e)

            return redirect('documentos:generar_enlace', doc_id=doc_id)
    else:
        form = EnlacePublicoForm(initial={'fecha_expiracion': documento.fecha_expiracion})

    enlace_generado = request.session.pop('enlace_generado', None)

    return render(request, 'documentos/generar_enlace.html', {
        'documento': documento,
        'form': form,
        'enlace_generado': enlace_generado,
    })


@login_required
def mis_enlaces(request):
    documentos = Documento.objects.filter(usuario=request.user).exclude(enlace_publico__isnull=True)
    return render(request, 'documentos/mis_enlaces.html', {'documentos': documentos, 'ahora': now()})


@require_POST
@login_required
def eliminar_enlace_publico(request, doc_id):
    try:
        documento = get_object_or_404(Documento, id=doc_id, usuario=request.user)
        documento.enlace_publico = None
        documento.fecha_expiracion = None
        documento.save()
        return JsonResponse({'success': True})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)